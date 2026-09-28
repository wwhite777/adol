"""Build the JKIICE-format Korean manuscript (docx) for KYRA-Bench phase A.
Result-independent text only: every outcome slot is a highlighted placeholder until the
CLAIM_LEDGER is filled from result/analysis/phaseA_T1/<stamp>/ (interim-look rule).
Citations: keys in ⟦...⟧ are numbered in order of first appearance."""
import json, re, pathlib, sys
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = pathlib.Path(__file__).resolve().parent
REVIEW = "--review" in sys.argv
ARGS = [a for a in sys.argv[1:] if a != "--review"]
OUT = HERE / (ARGS[0] if ARGS else ("kyra_phaseA_jkiice_ko_v5_review.docx" if REVIEW else "kyra_phaseA_jkiice_ko_v5.docx"))
DC = json.load(open(HERE / "refs_datacite_20260924.json", encoding="utf-8"))

KO_BODY, KO_HEAD, EN_SERIF, EN_SANS = "바탕", "맑은 고딕", "Times New Roman", "Arial"
MON = ["Jan.", "Feb.", "Mar.", "Apr.", "May", "Jun.", "Jul.", "Aug.", "Sep.", "Oct.", "Nov.", "Dec."]

# ---------------------------------------------------------------- references
def initials(name):
    if "," not in name:
        return name
    last, first = [x.strip() for x in name.split(",", 1)]
    parts = re.split(r"[\s]+", first)
    ini = []
    for p in parts:
        if not p:
            continue
        ini.append("-".join(q[0] + "." for q in p.split("-") if q))
    return " ".join(ini) + " " + last

def authors(names, team=None, max_n=10):
    names = [n for n in names if n not in (":", "Qwen")]
    if team:
        names = [n for n in names if n != team]
    fm = [initials(n) for n in names]
    if len(fm) > max_n:
        s = fm[0] + " et al."
    elif len(fm) == 1:
        s = fm[0]
    elif len(fm) == 2:
        s = fm[0] + " and " + fm[1]
    else:
        s = ", ".join(fm[:-1]) + ", and " + fm[-1]
    return (team + ", " + s) if team and fm else (team or s)

def arxiv_ref(aid, team=None):
    d = DC[aid]
    yy, mm = int(aid[:2]), int(aid[2:4])
    return (f"{authors(d['creators'], team)}, “{d['title']},” arXiv preprint arXiv: {aid}, "
            f"{MON[mm-1]} 20{yy:02d}. DOI: 10.48550/arXiv.{aid}.")

REFS = {
    "KIDBench": arxiv_ref("2605.25510"), "SproutBench": arxiv_ref("2508.11009"),
    "CAREBench": arxiv_ref("2606.29685"), "SafeChild": arxiv_ref("2506.13510"),
    "MinorBench": arxiv_ref("2503.10242"),
    "YouthSafe": ("Y. Yu, Y. Liu, Y. Zhang, Y. Huang, and Y. Wang, “YouthSafe: A youth-centric safety benchmark and safeguard "
                  "model for large language models,” in Proceeding of the 2025 ACM SIGSAC Conference on Computer and Communications "
                  "Security, Taipei, Taiwan, pp. 4349-4363, 2025. DOI: 10.1145/3719027.3765168."),
    "ChildSafe": arxiv_ref("2510.05484"), "Cha": arxiv_ref("2608.07902"),
    "RedQueen": ("Y. Jiang, K. Aggarwal, T. Laud, K. Munir, J. Pujara, and S. Mukherjee, “Red Queen: Exposing latent multi-turn "
                 "risks in large language models,” in Findings of the Association for Computational Linguistics: ACL 2025, Vienna, "
                 "Austria, pp. 25554-25591, 2025. DOI: 10.18653/v1/2025.findings-acl.1311."),
    "PersonaGrounded": ("P. Juneja and L. Lomidze, “Persona-grounded safety evaluation of AI companions in multi-turn conversations,” "
                        "in Proceeding of the 64th Annual Meeting of the Association for Computational Linguistics (Volume 1: Long "
                        "Papers), San Diego, CA, pp. 18148-18175, 2026. DOI: 10.18653/v1/2026.acl-long.828."),
    "CompanionBench": arxiv_ref("2608.02046"), "INTIMA": arxiv_ref("2508.09998"),
    "CogManip": arxiv_ref("2606.06099"), "TSJ": arxiv_ref("2606.25396"),
    "TAFMED": arxiv_ref("2608.10258"), "CultAdaptRT": arxiv_ref("2606.09178"),
    "ROKFORTRESS": arxiv_ref("2605.14152"), "KSAFEMM": arxiv_ref("2605.28013"),
    "Zheng": ("L. Zheng et al., “Judging LLM-as-a-judge with MT-Bench and Chatbot Arena,” in Advances in Neural Information "
              "Processing Systems 36, New Orleans, LA, pp. 46595-46623, 2023. DOI: 10.52202/075280-2020."), "NoFreeLabels": arxiv_ref("2503.05061"),
    "AICompanionBench": arxiv_ref("2606.04867"),
    "XSTest": ("P. Röttger, H. Kirk, B. Vidgen, G. Attanasio, F. Bianchi, and D. Hovy, “XSTest: A test suite for "
               "identifying exaggerated safety behaviours in large language models,” in Proceeding of the 2024 Conference "
               "of the North American Chapter of the Association for Computational Linguistics: Human Language "
               "Technologies (Volume 1: Long Papers), Mexico City, Mexico, pp. 5377-5400, 2024. "
               "DOI: 10.18653/v1/2024.naacl-long.301."),
    "Qwen25": arxiv_ref("2412.15115", team="Qwen Team"), "EXAONE4": arxiv_ref("2507.11407"),
    "Kanana": arxiv_ref("2502.18934", team="Kanana LLM Team"),
    "Gemma3": arxiv_ref("2503.19786", team="Gemma Team"),
    "HCX": arxiv_ref("2506.22403"),
    "HCXcard": ("NAVER Cloud. HyperCLOVAX-SEED-Think-14B [Internet]. Available: "
                "https://huggingface.co/naver-hyperclovax/HyperCLOVAX-SEED-Think-14B."),
    "KananaCard": ("Kakao Corp. kanana-1.5-8b-instruct-2505 [Internet]. Available: "
                   "https://huggingface.co/kakaocorp/kanana-1.5-8b-instruct-2505."),
    "Gwet2014": ("K. L. Gwet, Handbook of Inter-Rater Reliability, 4th ed. Gaithersburg, MD: Advanced Analytics, 2014."), "Llama3": arxiv_ref("2407.21783"), "Phi4": arxiv_ref("2412.08905"),
    "MistralSmall": ("Mistral AI. Mistral-Small-3.2-24B-Instruct-2506 [Internet]. Available: "
                     "https://huggingface.co/mistralai/Mistral-Small-3.2-24B-Instruct-2506."),
    "vLLM": ("W. Kwon, Z. Li, S. Zhuang, Y. Sheng, L. Zheng, C. H. Yu, J. Gonzalez, H. Zhang, and I. Stoica, "
             "“Efficient memory management for large language model serving with PagedAttention,” in Proceeding of "
             "the 29th Symposium on Operating Systems Principles, Koblenz, Germany, pp. 611-626, 2023. "
             "DOI: 10.1145/3600006.3613165."),
    "Gwet": ("K. L. Gwet, “Computing inter-rater reliability and its variance in the presence of high agreement,” "
             "British Journal of Mathematical and Statistical Psychology, vol. 61, no. 1, pp. 29-48, May 2008. "
             "DOI: 10.1348/000711006X126600."),
    "Holm": ("S. Holm, “A simple sequentially rejective multiple test procedure,” Scandinavian Journal of Statistics, "
             "vol. 6, no. 2, pp. 65-70, 1979. Available: https://www.jstor.org/stable/4615733."),
    "TOST": ("D. J. Schuirmann, “A comparison of the two one-sided tests procedure and the power approach for "
             "assessing the equivalence of average bioavailability,” Journal of Pharmacokinetics and "
             "Biopharmaceutics, vol. 15, no. 6, pp. 657-680, Dec. 1987. DOI: 10.1007/BF01068419."),
}

# ---------------------------------------------------------------- content
from content_v5 import *  # manuscript content v5 (review revision)

# ---------------------------------------------------------------- citation numbering
order = []
def scan(t):
    for m in re.finditer(r"⟦([^⟧]+)⟧", t):
        for k in m.group(1).split(","):
            k = k.strip()
            if k not in REFS:
                raise KeyError("unknown ref " + k)
            if k not in order:
                order.append(k)
for b in BODY:
    if b[0] in ("p", "ph"):
        scan(b[1])
def fmt_nums(nums):
    nums = sorted(nums); out = []; i = 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(f"{nums[i]}-{nums[j]}" if j - i >= 2 else ", ".join(str(n) for n in nums[i:j + 1]))
        i = j + 1
    return "[" + ", ".join(out) + "]"
def cite(t):
    return re.sub(r"⟦([^⟧]+)⟧", lambda m: fmt_nums([order.index(k.strip()) + 1 for k in m.group(1).split(",")]), t)
unused = [k for k in REFS if k not in order]

# ---------------------------------------------------------------- docx helpers
doc = Document()
def set_fonts(run, latin, east, size, bold=False, italic=False):
    run.font.name = latin; run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), latin)
    rf.set(qn("w:eastAsia"), east)

def para(text="", latin=EN_SERIF, east=KO_BODY, size=9.2, bold=False, align=None, indent_pt=None,
         spacing=1.5, before=0, after=0, italic=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = spacing; pf.space_before = Pt(before); pf.space_after = Pt(after)
    if indent_pt is not None:
        pf.first_line_indent = Pt(indent_pt)
    if align:
        p.alignment = align
    add_runs(p, text, latin, east, size, bold, italic)
    return p

def add_runs(p, text, latin, east, size, bold=False, italic=False):
    # ⟪...⟫ = result placeholder, rendered highlighted
    for i, seg in enumerate(re.split(r"(⟪[^⟫]*⟫)", text)):
        if not seg:
            continue
        r = p.add_run(seg)
        set_fonts(r, latin, east, size, bold, italic)
        if seg.startswith("⟪"):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW

def set_cols(section, n, space_mm=6):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols"); sectPr.append(cols)
    cols.set(qn("w:num"), str(n)); cols.set(qn("w:space"), str(int(space_mm / 25.4 * 1440)))

def page_setup(s):
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.top_margin = s.bottom_margin = Mm(38.5); s.left_margin = s.right_margin = Mm(30)
    s.header_distance = s.footer_distance = Mm(10)

def cell_text(cell, text, bold=False, size=8):
    cell.text = ""
    p = cell.paragraphs[0]; p.paragraph_format.line_spacing = 1.2
    add_runs(p, text, EN_SERIF, KO_BODY, size, bold)

def set_cell_border(table):
    tbl = table._tbl; tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000")
        b.append(e)
    for edge in ("left", "right", "insideV"):
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "nil"); b.append(e)
    tblPr.append(b)

def add_table(key):
    cap, head, rows, bold_last = TABLES[key]
    para(cap, EN_SANS, KO_HEAD, 8, spacing=1.3, before=4, after=2)
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_border(t)
    for j, h in enumerate(head):
        cell_text(t.rows[0].cells[j], h, bold=True)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            cell_text(t.rows[i + 1].cells[j], v, bold=(bold_last and i == len(rows) - 1))
    para("", size=4)

def one_col_block(fn):
    s = doc.add_section(WD_SECTION.CONTINUOUS); page_setup(s); set_cols(s, 1)
    fn()
    s2 = doc.add_section(WD_SECTION.CONTINUOUS); page_setup(s2); set_cols(s2, 2)

def add_figure(key):
    fname, cap = FIGS[key]
    def f():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        from PIL import Image
        jpg = HERE / (pathlib.Path(fname).stem + ".jpg")
        Image.open(HERE / fname).convert("RGB").save(jpg, "JPEG", quality=95, dpi=(300, 300))
        p.add_run().add_picture(str(jpg), width=Mm(95 if key == "fig3" else 145))
        para(cap, EN_SANS, KO_HEAD, 8, spacing=1.3, after=6)
    one_col_block(f)

# ---------------------------------------------------------------- front matter (1 column)
s0 = doc.sections[0]; page_setup(s0); set_cols(s0, 1)
para(TITLE_KO, EN_SANS, KO_HEAD, 17, bold=True, spacing=1.3, after=6)
if not REVIEW:
    p = para("", spacing=1.2, after=6)
    add_runs(p, "정원철", EN_SERIF, KO_BODY, 12); r = p.add_run("1"); set_fonts(r, EN_SERIF, KO_BODY, 8); r.font.superscript = True
    add_runs(p, " · 오하영", EN_SERIF, KO_BODY, 12); r = p.add_run("2*"); set_fonts(r, EN_SERIF, KO_BODY, 8); r.font.superscript = True
para(TITLE_EN, EN_SANS, KO_HEAD, 12.5, bold=True, spacing=1.2, after=6)
if not REVIEW:
    p = para("", spacing=1.2, after=4)
    add_runs(p, "Woncheol Jeong", EN_SERIF, KO_BODY, 10.5); r = p.add_run("1"); set_fonts(r, EN_SERIF, KO_BODY, 7); r.font.superscript = True
    add_runs(p, " · Hayoung Oh", EN_SERIF, KO_BODY, 10.5); r = p.add_run("2*"); set_fonts(r, EN_SERIF, KO_BODY, 7); r.font.superscript = True
    for aff in ["1Research Professor, Intelligent Digital Therapeutics Research Center, Sungkyunkwan University, Seoul, 03063 Korea",
                "2Associate Professor, Department of Applied AI, Sungkyunkwan University, Seoul, 03063 Korea",
                "*Corresponding author"]:
        p = para("", spacing=1.2)
        if aff[0].isdigit():
            r = p.add_run(aff[0]); set_fonts(r, EN_SERIF, KO_BODY, 7); r.font.superscript = True
            add_runs(p, aff[1:], EN_SERIF, KO_BODY, 9)
        else:
            add_runs(p, aff, EN_SERIF, KO_BODY, 9)
para("요  약", EN_SANS, KO_HEAD, 10, bold=True, spacing=1.6, before=10)
para(cite(ABS_KO), EN_SERIF, KO_BODY, 9.2, indent_pt=10, spacing=1.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
para("ABSTRACT", EN_SANS, KO_HEAD, 10, bold=True, spacing=1.6, before=8)
para(ABS_EN, EN_SERIF, KO_BODY, 8.5, indent_pt=10, spacing=1.3, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
p = para("", spacing=1.56, before=8)
add_runs(p, "키워드 : ", EN_SANS, KO_HEAD, 10, bold=True); add_runs(p, KW_KO, EN_SERIF, KO_BODY, 8.5)
p = para("", spacing=1.56)
add_runs(p, "Keywords : ", EN_SANS, KO_HEAD, 10, bold=True); add_runs(p, KW_EN, EN_SERIF, KO_BODY, 8.5)
if not REVIEW:
    para("* Corresponding Author Hayoung Oh (E-mail: hyoh79@skku.edu, Tel: +82-10-5389-5996)", EN_SERIF, KO_BODY, 8, spacing=1.2, before=14)
    para("Associate Professor, Department of Applied AI, Sungkyunkwan University, Seoul, 03063 Korea", EN_SERIF, KO_BODY, 8, spacing=1.2)

# ---------------------------------------------------------------- body (2 columns, new page)
sb = doc.add_section(WD_SECTION.NEW_PAGE); page_setup(sb); set_cols(sb, 2)
BODY_KW = dict(latin=EN_SERIF, east=KO_BODY, size=9.2, spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
first_h1 = True
for b in BODY:
    kind = b[0]
    if kind == "h1":
        para(b[1], EN_SANS, KO_HEAD, 11, bold=True, spacing=1.5, before=0 if first_h1 else 14, after=6)
        first_h1 = False
    elif kind == "h2":
        para(b[1], EN_SANS, KO_HEAD, 9.5, bold=True, spacing=1.5, before=6, after=2)
    elif kind in ("p", "ph"):
        para(cite(b[1]), indent_pt=9.2 * 1, **BODY_KW)
    elif kind == "eq":
        p = para(b[1] + "        " + b[2], EN_SERIF, KO_BODY, 9.2, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.5, before=4, after=4, italic=False)
    elif kind == "table":
        if TABLES[b[1]][1] and len(TABLES[b[1]][1]) >= 4 and b[1] in ("t2", "t3", "t_endpoints", "t_chrono", "t4", "t5", "t_validity"):
            one_col_block(lambda k=b[1]: add_table(k))
        else:
            add_table(b[1])
    elif kind == "figure":
        add_figure(b[1])

if not REVIEW:
    para("ACKNOWLEDGEMENTS", EN_SANS, KO_HEAD, 11, bold=True, spacing=1.5, before=10, after=2)
    para(ACK, EN_SERIF, KO_BODY, 9.2, spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
para("생성형 AI 활용 공개", EN_SANS, KO_HEAD, 9.5, bold=True, spacing=1.5, before=6, after=2)
para(AI_USE, indent_pt=9.2, **BODY_KW)
para("REFERENCES", EN_SANS, KO_HEAD, 11, bold=True, spacing=1.5, before=14, after=4)
for i, k in enumerate(order, 1):
    p = para(f"[{i}] " + REFS[k], EN_SERIF, KO_BODY, 8, spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    p.paragraph_format.left_indent = Pt(16); p.paragraph_format.first_line_indent = Pt(-16)

# author bios (as in the authors' previous JKIICE paper)
BIOS = [("정원철(Woncheol Jeong)", ["1999년 2월 : 서울대학교 재료공학부(공학사)", "2001년 8월 : 서울대학교 재료공학부(공학석사)",
         "2005년 2월 : MIT 기계공학 생명공학 (공학석사)", "2017년 2월 : 서울대학교 재료공학부(공학박사)",
         "2018년 4월 ～ 2018년 6월: Beihang Univ, 연구교수", "2018년 10월 ～ 2019년 4월: Ton Duc Thang Univ, 선임연구원",
         "2019년 9월 ～ 2020년 9월: Far Eastern Federal Univ, 부교수", "2021년 9월 ～ 2022년 8월: 서울시립대학교, 강사",
         "2022년 3월 ～ 2023년 2월: 고려대학교, 연구교수", "2023년 3월 ～ 2023년 12월: 포항공과대학교, 연구교수",
         "2024년 1월 ～ 2025년 1월: 난양공과대학교, 책임연구원", "2024년 9월 ～ 현재: 성균관대학교, 연구교수",
         "※관심분야 : AI, 생명공학, 재료, 기계, Quantum"]),
        ("오하영(Hayoung Oh)", ["2013년 2월 : 서울대학교 컴퓨터공학부(공학박사)", "2013년 3월 ～ 2013년 8월: 서울시립대학교, 연구교수",
         "2013년 9월 ～ 2016년 8월: 숭실대학교, 조교수", "2016년 9월 ～ 2020년 2월: 아주대학교, 조교수",
         "2020년 3월 ～ 현재: 성균관대학교, 부교수", "2025년 3월 ～ 현재: 인공지능융합학과, 학과장",
         "2024년 4월 ～ 현재: 지능형디지털치료기기센터, 센터장", "※관심분야 : AI"])]
for name, lines in ([] if REVIEW else BIOS):
    para(name, EN_SANS, KO_HEAD, 9, bold=True, spacing=1.3, before=10)
    para("약력", EN_SERIF, KO_BODY, 8, spacing=1.3)
    for ln in lines:
        para(("" if ln.startswith("※") else "․") + ln, EN_SERIF, KO_BODY, 8, spacing=1.3)

cp = doc.core_properties
cp.author = "" if REVIEW else "Woncheol Jeong; Hayoung Oh"
cp.last_modified_by = ""; cp.comments = ""; cp.title = TITLE_EN; cp.keywords = ""
doc.save(OUT)
print("wrote", OUT)
print("references cited:", len(order), "unused:", unused)
