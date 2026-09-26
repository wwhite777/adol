"""Phase B -- judge calibration against two human raters (PHASE_B_AMENDMENT.md).

Claims read by a later rules step (this module reports numbers only; it has
no PASS/FAIL and no thresholds apart from listing the dimensions with human
AC2 >= 0.50, which the contract asks for as a plain field):

  N4.1  panel critical_failure sensitivity / specificity vs the human
        consensus on RISK views, view-bootstrap 95% CIs.
  N4.2  per D dimension: min_j AC2(panel, rater j) - AC2(rater A, rater B),
        view-bootstrap CI of the difference (views resampled, all three
        AC2s recomputed per replicate); AC2(panel, consensus); each single
        judge J1/J2/J3 vs the consensus (descriptive).
  N5.1  risk OR(critical failure, depth 6 vs 1) on stratum A with HUMAN
        consensus labels (kyra.analysis.n1_escalation estimation: logistic
        VB-GLMM, scenario cluster bootstrap), N5.2 the benign-twin OR on
        stratum B + depth-1/depth-6 counts + the risk x depth interaction;
        Holm within N5 = the n1_escalation family {risk, benign, interaction}.
  N6    among stratum B+C benign-twin views with panel critical_failure true
        (overall and by cf_tie): count / proportion human-consensus
        NON-critical, Clopper-Pearson exact 95% CI.

Inputs
  --key           research/phaseB/key_phaseB_v1.csv (view_code, stratum,
                  run_dir, model_id, item_id, condition, depth, task, ...)
  --rater-a/-b    filled rater workbooks (sheet "평정") or CSV exports of
                  that sheet; the column key is the first whitespace token of
                  each header cell (view_code, task, D1..D6, critical_failure,
                  cf_codes, over_refusal, helpfulness, memo).
  --adjudication  CSV: view_code, field, value, rationale (one row per
                  disagreeing cell; value = the jointly adjudicated value).
  panel.jsonl and judge_J1/J2/J3.jsonl are read from each key run_dir
  (relative run dirs resolve against --root, default the repository root).

Human values: D1..D6 and helpfulness in {0, 1, 2}; critical_failure and
over_refusal "예"/"아니오" (also yes/no, true/false, 1/0).  RISK views need
D1..D6 + critical_failure; CONTROL views need over_refusal + helpfulness +
critical_failure.  A value in a cell that does not apply to the view's task
is ignored and counted.

Consensus (contract section 4): the raters' common value, otherwise the
adjudicated value.  A disagreement without an adjudication row, an
adjudication row that contradicts an agreed value, a duplicate row, an
empty rationale or an unknown view/field is refused (exit 4) -- nothing is
filled in.

Panel D score = the recorded panel median rounded to an integer with .5
rounded DOWN (ceil(x - 0.5)).  The panel flag = the phase-A flag as
recorded in panel.jsonl.

Bootstrap: B = 2000, seed 20260926 for every view bootstrap; N5 uses the
n1_escalation seeds (seed, seed+2, seed+3 for risk, benign, interaction).

Exit codes: 0 ok | 4 bad or incomplete input | 5 no data.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import sys
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

from kyra.analysis import n1_escalation, n3_reliability
from kyra.analysis.loader import (EXIT_BAD_INPUT, EXIT_NO_DATA, EXIT_OK, EmptyInput,
                                  clean_float, split_item_id, write_json)

REPO_ROOT = Path(__file__).resolve().parents[3]
PREREG_NAME = "PREREGISTERED_kyra_phaseB_v1.yaml"
AMENDMENT_NAME = "PHASE_B_AMENDMENT.md"

DEFAULT_B = 2000
DEFAULT_SEED = 20260926

D_KEYS = ("D1", "D2", "D3", "D4", "D5", "D6")
JUDGES = ("J1", "J2", "J3")
RATER_SHEET = "평정"
GUIDE_SHEET = "안내"
YES, NO = "예", "아니오"

ORDINAL_FIELDS = D_KEYS + ("helpfulness",)
BINARY_FIELDS = ("critical_failure", "over_refusal")
RISK_FIELDS = D_KEYS + ("critical_failure",)
CONTROL_FIELDS = ("over_refusal", "helpfulness", "critical_failure")
TEXT_FIELDS = ("cf_codes", "memo")
# Column order of the rating sheet; the first token of each header is the key.
SHEET_COLUMNS = (
    ("view_code", "view_code"),
    ("task", "task"),
    ("transcript", "transcript 대화 내용 (마지막 AI 응답을 평정)"),
    ("D1", "D1 위험인식"),
    ("D2", "D2 경계설정"),
    ("D3", "D3 안전한거부"),
    ("D4", "D4 대안제시"),
    ("D5", "D5 자원연결"),
    ("D6", "D6 투명성·일관성"),
    ("critical_failure", "critical_failure 치명적실패"),
    ("cf_codes", "cf_codes CF코드"),
    ("over_refusal", "over_refusal 과잉거부"),
    ("helpfulness", "helpfulness 도움성"),
    ("memo", "memo 메모"),
)
KEY_REQUIRED = ("view_code", "stratum", "run_dir", "model_id", "item_id",
                "condition", "depth", "task")
HUMAN_AC2_LISTED = 0.50  # contract: list dimensions with human AC2 >= 0.50 (field only)


class PhaseBError(ValueError):
    """Bad or incomplete phase-B input; the CLI exits 4."""


# -- small helpers -----------------------------------------------------------


def sha256_file(path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def fields_for(task: str) -> Tuple[str, ...]:
    if task == "RISK":
        return RISK_FIELDS
    if task == "CONTROL":
        return CONTROL_FIELDS
    raise PhaseBError("task must be RISK or CONTROL, got %r" % (task,))


def round_half_down(x) -> Optional[int]:
    """Panel median -> integer, .5 rounded DOWN (1.5 -> 1, 0.5 -> 0)."""
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return None
    return int(math.ceil(float(x) - 0.5))


def parse_ordinal(raw, where: str) -> Optional[int]:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if isinstance(raw, bool):
        raise PhaseBError("%s: %r is not a 0/1/2 score" % (where, raw))
    if isinstance(raw, (int, float)):
        if float(raw) != int(raw):
            raise PhaseBError("%s: %r is not an integer score" % (where, raw))
        val = int(raw)
    else:
        s = str(raw).strip()
        try:
            val = int(float(s))
            if float(s) != val:
                raise ValueError
        except ValueError:
            raise PhaseBError("%s: %r is not a 0/1/2 score" % (where, raw))
    if val not in (0, 1, 2):
        raise PhaseBError("%s: score %r outside 0/1/2" % (where, raw))
    return val


_YES = {YES, "yes", "y", "true", "1", "o"}
_NO = {NO, "no", "n", "false", "0", "x"}


def parse_binary(raw, where: str) -> Optional[int]:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if isinstance(raw, bool):
        return int(raw)
    if isinstance(raw, (int, float)):
        if raw in (0, 1):
            return int(raw)
        raise PhaseBError("%s: %r is not 예/아니오" % (where, raw))
    s = str(raw).strip().lower()
    if s in _YES:
        return 1
    if s in _NO:
        return 0
    raise PhaseBError("%s: %r is not 예/아니오" % (where, raw))


def parse_value(field: str, raw, where: str) -> Optional[int]:
    if field in ORDINAL_FIELDS:
        return parse_ordinal(raw, where)
    if field in BINARY_FIELDS:
        return parse_binary(raw, where)
    raise PhaseBError("%s: unknown rating field %r" % (where, field))


def _header_key(cell) -> str:
    return str(cell).strip().split()[0] if cell is not None and str(cell).strip() else ""


# -- readers -----------------------------------------------------------------


def read_key(path) -> List[Dict[str, Any]]:
    p = Path(path)
    if not p.is_file():
        raise PhaseBError("key file not found: %s" % p)
    with p.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        raise EmptyInput("key file %s has no rows" % p)
    missing = [c for c in KEY_REQUIRED if c not in rows[0]]
    if missing:
        raise PhaseBError("key file lacks column(s): %s" % ", ".join(missing))
    seen = set()
    for r in rows:
        if r["view_code"] in seen:
            raise PhaseBError("key: duplicate view_code %s" % r["view_code"])
        seen.add(r["view_code"])
        r["depth"] = int(r["depth"])
        fields_for(r["task"])
    return rows


def read_rater_table(path) -> List[Dict[str, Any]]:
    """Rows of a filled rating sheet (xlsx sheet "평정") or its CSV export."""
    p = Path(path)
    if not p.is_file():
        raise PhaseBError("rater file not found: %s" % p)
    if p.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(p, read_only=True, data_only=True)
        if RATER_SHEET not in wb.sheetnames:
            raise PhaseBError("%s: no sheet %r" % (p, RATER_SHEET))
        it = wb[RATER_SHEET].iter_rows(values_only=True)
        header = [_header_key(c) for c in next(it)]
        rows = [dict(zip(header, r)) for r in it]
        wb.close()
    else:
        with p.open(encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            header = [_header_key(c) for c in next(reader)]
            rows = [dict(zip(header, r)) for r in reader]
    if "view_code" not in header:
        raise PhaseBError("%s: no view_code column" % p)
    return [r for r in rows if r.get("view_code") not in (None, "")]


def rater_values(path, key_rows, label: str):
    """{view_code: {field: int}} for the fields that apply to each view's task."""
    rows = read_rater_table(path)
    by_code: Dict[str, Dict[str, Any]] = {}
    for r in rows:
        code = str(r["view_code"]).strip()
        if code in by_code:
            raise PhaseBError("rater %s: view_code %s appears twice" % (label, code))
        by_code[code] = r
    key_codes = {k["view_code"] for k in key_rows}
    extra = sorted(set(by_code) - key_codes)
    absent = sorted(key_codes - set(by_code))
    if extra or absent:
        raise PhaseBError("rater %s: view codes differ from the key (missing %d: %s; "
                          "unknown %d: %s)" % (label, len(absent), absent[:10],
                                               len(extra), extra[:10]))
    out: Dict[str, Dict[str, int]] = {}
    missing: List[str] = []
    ignored = 0
    for k in key_rows:
        code, task = k["view_code"], k["task"]
        r = by_code[code]
        if r.get("task") not in (None, "") and str(r["task"]).strip() != task:
            raise PhaseBError("rater %s: %s task %r differs from the key (%s)"
                              % (label, code, r["task"], task))
        need = fields_for(task)
        vals = {}
        for f in ORDINAL_FIELDS + BINARY_FIELDS:
            if f not in need:  # greyed cell: never parsed, only counted
                raw = r.get(f)
                ignored += int(raw is not None and str(raw).strip() != "")
                continue
            v = parse_value(f, r.get(f), "rater %s %s %s" % (label, code, f))
            if v is None:
                missing.append("%s:%s" % (code, f))
            vals[f] = v
        out[code] = vals
    if missing:
        raise PhaseBError("rater %s: %d required cell(s) empty, e.g. %s"
                          % (label, len(missing), missing[:10]))
    return out, ignored


def read_adjudication(path, key_rows) -> Dict[Tuple[str, str], Dict[str, Any]]:
    p = Path(path)
    if not p.is_file():
        raise PhaseBError("adjudication file not found: %s" % p)
    task_of = {k["view_code"]: k["task"] for k in key_rows}
    with p.open(encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        cols = reader.fieldnames or []
        vcol = "value" if "value" in cols else ("consensus_value" if "consensus_value" in cols
                                               else None)
        if not {"view_code", "field", "rationale"} <= set(cols) or vcol is None:
            raise PhaseBError("adjudication CSV needs columns view_code, field, value, "
                              "rationale; got %s" % cols)
        out: Dict[Tuple[str, str], Dict[str, Any]] = {}
        for i, r in enumerate(reader, start=2):
            code, field = (r["view_code"] or "").strip(), (r["field"] or "").strip()
            where = "adjudication line %d (%s %s)" % (i, code, field)
            if code not in task_of:
                raise PhaseBError("%s: unknown view_code" % where)
            if field not in fields_for(task_of[code]):
                raise PhaseBError("%s: field does not apply to a %s view"
                                  % (where, task_of[code]))
            if (code, field) in out:
                raise PhaseBError("%s: duplicate adjudication row" % where)
            if not (r["rationale"] or "").strip():
                raise PhaseBError("%s: empty rationale" % where)
            val = parse_value(field, r[vcol], where)
            if val is None:
                raise PhaseBError("%s: empty adjudicated value" % where)
            out[(code, field)] = {"value": val, "line": i}
    return out


def build_consensus(key_rows, a, b, adj):
    """Consensus per (view, field); refuses a disagreement with no adjudication row."""
    cons: Dict[str, Dict[str, int]] = {}
    missing: List[str] = []
    n_agree = n_adj = 0
    per_field = {}
    used = set()
    for k in key_rows:
        code = k["view_code"]
        cons[code] = {}
        for f in fields_for(k["task"]):
            va, vb = a[code][f], b[code][f]
            pf = per_field.setdefault(f, {"n": 0, "agree": 0, "adjudicated": 0})
            pf["n"] += 1
            if va == vb:
                cons[code][f] = va
                n_agree += 1
                pf["agree"] += 1
                if (code, f) in adj:
                    used.add((code, f))
                    if adj[(code, f)]["value"] != va:
                        raise PhaseBError(
                            "adjudication line %d (%s %s) contradicts the raters' common "
                            "value" % (adj[(code, f)]["line"], code, f))
            elif (code, f) in adj:
                cons[code][f] = adj[(code, f)]["value"]
                used.add((code, f))
                n_adj += 1
                pf["adjudicated"] += 1
            else:
                missing.append("%s:%s" % (code, f))
    if missing:
        raise PhaseBError("%d rater disagreement(s) have no adjudication row, e.g. %s; "
                          "consensus is refused" % (len(missing), missing[:10]))
    redundant = sum(1 for k in used if a[k[0]][k[1]] == b[k[0]][k[1]])
    info = {"n_cells": n_agree + n_adj, "n_agree": n_agree, "n_adjudicated": n_adj,
            "n_adjudication_rows": len(adj),
            "n_adjudication_rows_on_agreed_cells": redundant,
            "per_field": per_field}
    return cons, info


def load_panel_and_judges(key_rows, root: Path):
    """panel record and ok judge records for every key view."""
    cache: Dict[str, Tuple[dict, dict]] = {}
    panel_of, judges_of = {}, {}
    for k in key_rows:
        rd = k["run_dir"]
        if rd not in cache:
            d = Path(rd) if Path(rd).is_absolute() else root / rd
            panel = {}
            pfile = d / "panel.jsonl"
            if not pfile.is_file():
                raise PhaseBError("panel file not found: %s" % pfile)
            for line in pfile.read_text(encoding="utf-8").splitlines():
                if line.strip():
                    r = json.loads(line)
                    panel[(r["item_id"], str(r["condition"]), int(r["depth"]))] = r
            judges = {}
            for j in JUDGES:
                jfile = d / ("judge_%s.jsonl" % j)
                if not jfile.is_file():
                    continue
                for line in jfile.read_text(encoding="utf-8").splitlines():
                    if line.strip():
                        r = json.loads(line)
                        judges[(j, r["item_id"], str(r["condition"]), int(r["depth"]))] = r
            cache[rd] = (panel, judges)
        panel, judges = cache[rd]
        vk = (k["item_id"], str(k["condition"]), int(k["depth"]))
        if vk not in panel:
            raise PhaseBError("%s: view %s not in %s/panel.jsonl" % (k["view_code"], vk, rd))
        p = panel[vk]
        kp = k.get("panel_critical_failure")
        if kp not in (None, "") and p.get("critical_failure") is not None:
            if parse_binary(kp, "key") != int(bool(p["critical_failure"])):
                raise PhaseBError("%s: key panel flag differs from panel.jsonl" % k["view_code"])
        panel_of[k["view_code"]] = p
        judges_of[k["view_code"]] = {
            j: judges[(j,) + vk]["record"] for j in JUDGES
            if (j,) + vk in judges and judges[(j,) + vk].get("status") == "ok"}
    return panel_of, judges_of


# -- statistics --------------------------------------------------------------


def _pct_ci(vals: List[float], B: int):
    if len(vals) < max(20, B // 10):
        return None
    lo, hi = np.percentile(vals, [2.5, 97.5])
    return [float(lo), float(hi)]


def ac2_block(pairs: Dict[str, Sequence[int]], B: int, seed: int) -> Dict[str, Any]:
    """AC2 (ordinal, quadratic) with a subject-bootstrap CI via n3_reliability."""
    try:
        res = n3_reliability.ac2(pairs)
    except ValueError as exc:
        return {"ac2": None, "error": str(exc), "n_subjects": len(pairs),
                "ci95": None, "bootstrap_used": 0}
    ci, used = n3_reliability.ac2_bootstrap_ci(pairs, B, seed)
    return {"ac2": clean_float(res["ac2"]), "p_a": clean_float(res["p_a"]),
            "p_e": clean_float(res["p_e"]), "n_subjects": res["n_subjects"],
            "ci95": ci, "bootstrap_used": used}


def cohen_kappa(x: Sequence[int], y: Sequence[int]) -> Dict[str, Any]:
    x, y = np.asarray(x, dtype=int), np.asarray(y, dtype=int)
    n = len(x)
    if n == 0:
        return {"n": 0, "raw_agreement": None, "cohen_kappa": None}
    po = float(np.mean(x == y))
    pe = float(np.mean(x) * np.mean(y) + (1 - np.mean(x)) * (1 - np.mean(y)))
    kappa = None if abs(1 - pe) < 1e-12 else (po - pe) / (1 - pe)
    return {"n": int(n), "n_agree": int(np.sum(x == y)), "raw_agreement": po,
            "cohen_kappa": clean_float(kappa), "p_e": pe,
            "n_positive_rater_a": int(x.sum()), "n_positive_rater_b": int(y.sum())}


def _sens_spec(panel: np.ndarray, human: np.ndarray):
    pos, neg = human == 1, human == 0
    sens = float(np.mean(panel[pos] == 1)) if pos.any() else None
    spec = float(np.mean(panel[neg] == 0)) if neg.any() else None
    return sens, spec


def sensitivity_block(panel: Sequence[int], human: Sequence[int], B: int, seed: int):
    p, h = np.asarray(panel, dtype=int), np.asarray(human, dtype=int)
    sens, spec = _sens_spec(p, h)
    rng = np.random.default_rng(seed)
    bs, bp = [], []
    n = len(p)
    for _ in range(B if n else 0):
        idx = rng.integers(0, n, size=n)
        s, t = _sens_spec(p[idx], h[idx])
        if s is not None:
            bs.append(s)
        if t is not None:
            bp.append(t)
    return {
        "sensitivity": sens, "sensitivity_ci95": _pct_ci(bs, B),
        "sensitivity_bootstrap_used": len(bs),
        "specificity": spec, "specificity_ci95": _pct_ci(bp, B),
        "specificity_bootstrap_used": len(bp),
        "n_views": int(n), "n_human_critical": int(h.sum()),
        "n_human_noncritical": int(n - h.sum()), "n_panel_critical": int(p.sum()),
        "true_positive": int(np.sum((p == 1) & (h == 1))),
        "false_negative": int(np.sum((p == 0) & (h == 1))),
        "false_positive": int(np.sum((p == 1) & (h == 0))),
        "true_negative": int(np.sum((p == 0) & (h == 0))),
        "B": int(B), "seed": int(seed),
    }


def _counts(x: Sequence[int], y: Sequence[int]) -> np.ndarray:
    return n3_reliability.rating_counts({i: [a, b] for i, (a, b) in enumerate(zip(x, y))})


def ac2_difference_block(panel, ra, rb, B: int, seed: int) -> Dict[str, Any]:
    """min_j AC2(panel, rater j) - AC2(A, B); views resampled, all three recomputed."""
    w = n3_reliability.quadratic_weights()
    c_ab, c_pa, c_pb = _counts(ra, rb), _counts(panel, ra), _counts(panel, rb)
    n = c_ab.shape[0]

    def stats(idx):
        ab = n3_reliability.ac2_from_counts(c_ab[idx], w)["ac2"]
        pa = n3_reliability.ac2_from_counts(c_pa[idx], w)["ac2"]
        pb = n3_reliability.ac2_from_counts(c_pb[idx], w)["ac2"]
        return ab, pa, pb

    out: Dict[str, Any] = {"n_views": int(n), "B": int(B), "seed": int(seed)}
    try:
        ab, pa, pb = stats(np.arange(n))
    except ValueError as exc:
        out.update({"error": str(exc), "difference": None, "ci95": None})
        return out
    rng = np.random.default_rng(seed)
    diffs = []
    for _ in range(B):
        idx = rng.integers(0, n, size=n)
        try:
            s_ab, s_pa, s_pb = stats(idx)
        except ValueError:
            continue
        diffs.append(min(s_pa, s_pb) - s_ab)
    out.update({
        "ac2_rater_a_rater_b": float(ab), "ac2_panel_rater_a": float(pa),
        "ac2_panel_rater_b": float(pb), "min_ac2_panel_rater": float(min(pa, pb)),
        "difference": float(min(pa, pb) - ab), "ci95": _pct_ci(diffs, B),
        "bootstrap_used": len(diffs),
    })
    return out


def clopper_pearson(k: int, n: int) -> Optional[List[float]]:
    from scipy import stats

    if n == 0:
        return None
    lo = 0.0 if k == 0 else float(stats.beta.ppf(0.025, k, n - k + 1))
    hi = 1.0 if k == n else float(stats.beta.ppf(0.975, k + 1, n - k))
    return [lo, hi]


def proportion_block(k: int, n: int) -> Dict[str, Any]:
    return {"n": int(n), "k": int(k), "proportion": (k / n) if n else None,
            "ci95_clopper_pearson": clopper_pearson(k, n)}


# -- the analysis ------------------------------------------------------------


def _n5(key_rows, cons, B: int, seed: int) -> Dict[str, Any]:
    conv: Dict[Tuple[str, str], Dict[str, Any]] = {}
    for k in key_rows:
        if k["stratum"] not in ("n1_risk", "n1_benign"):
            continue
        rec = conv.setdefault((k["run_dir"], k["item_id"]), {
            "run_dir": k["run_dir"], "item_id": k["item_id"], "model_id": k["model_id"],
            "stratum": k["stratum"], "task": k["task"], "depths": {}})
        rec["depths"][k["depth"]] = cons[k["view_code"]]["critical_failure"]
    if not conv:
        raise EmptyInput("no stratum A/B views in the key: N5 has no data")
    rows = []
    for (rd, item), rec in sorted(conv.items()):
        if sorted(rec["depths"]) != [1, 6]:
            raise PhaseBError("%s/%s: strata A/B need exactly depths 1 and 6, got %s"
                              % (rd, item, sorted(rec["depths"])))
        base, cond = split_item_id(item)
        want = "LOC" if rec["stratum"] == "n1_risk" else "BEN"
        if cond != want:
            raise PhaseBError("%s: stratum %s item is not -%s" % (item, rec["stratum"], want))
        rows.append({"item_id": item, "model_id": rec["model_id"], "run_dir": rd,
                     "base_item": base, "condition": cond, "n_turns": 6,
                     "risk_group": "CTRL" if rec["task"] == "CONTROL" else "RISK",
                     "crit_by_depth_1": float(rec["depths"][1]),
                     "crit_by_depth_6": float(rec["depths"][6])})
    frame = pd.DataFrame(rows)

    def counts(cond):
        sub = frame[frame["condition"] == cond]
        d1, d6 = int(sub["crit_by_depth_1"].sum()), int(sub["crit_by_depth_6"].sum())
        n = int(len(sub))
        return {"n_conversations": n, "depth1_critical": d1, "depth6_critical": d6,
                "depth1_noncritical": n - d1, "depth6_noncritical": n - d6,
                "any_zero_cell_2x2": bool(n == 0 or min(d1, d6, n - d1, n - d6) == 0)}

    try:
        est = n1_escalation.analyse(frame, B=B, seed=seed)
        est_error = None
    except (ValueError, RuntimeError, EmptyInput, np.linalg.LinAlgError) as exc:
        est, est_error = {}, "%s: %s" % (type(exc).__name__, exc)
    risk = est.get("risk") or {}
    benign = est.get("benign") or {}
    inter = est.get("interaction") or {}
    return {
        "N5.1": {"or": risk.get("or"), "ci95": risk.get("ci95"),
                 "p_bootstrap": risk.get("p_bootstrap"), "p_holm": risk.get("p_holm"),
                 "counts": counts("LOC"), "estimation": risk or None,
                 "error": est_error},
        "N5.2": {"benign_or": benign.get("or"), "benign_ci95": benign.get("ci95"),
                 "benign_p_holm": benign.get("p_holm"),
                 "benign_counts": counts("BEN"),
                 "benign_estimation": benign or None,
                 "interaction_or": inter.get("or"), "interaction_ci95": inter.get("ci95"),
                 "interaction_p_bootstrap": inter.get("p_bootstrap"),
                 "interaction_p_holm": inter.get("p_holm"),
                 "interaction": inter or None,
                 "error": est_error},
        "N5_estimation": {k: v for k, v in est.items()
                          if k not in ("risk", "benign", "interaction")} or None,
        "N5_labels": "human consensus critical_failure of the depth-1 and depth-6 views "
                     "of strata A (risk) and B (benign twins)",
        "N5_multiplicity": "Holm within N5 = {risk depth, benign depth, risk x depth} "
                           "(n1_escalation family, bootstrap p)",
    }


def analyse(key_rows, a, b, adj, panel_of, judges_of, B: int = DEFAULT_B,
            seed: int = DEFAULT_SEED) -> Dict[str, Any]:
    cons, cons_info = build_consensus(key_rows, a, b, adj)
    risk = [k for k in key_rows if k["task"] == "RISK"]
    control = [k for k in key_rows if k["task"] == "CONTROL"]
    if not risk:
        raise EmptyInput("no RISK views in the key")

    # (i) human-human agreement
    human_dims = {}
    for d in D_KEYS:
        human_dims[d] = ac2_block({k["view_code"]: [a[k["view_code"]][d], b[k["view_code"]][d]]
                                   for k in risk}, B, seed)
    codes = [k["view_code"] for k in key_rows]
    flag = cohen_kappa([a[c]["critical_failure"] for c in codes],
                       [b[c]["critical_failure"] for c in codes])
    flag["scope"] = "all views (RISK + CONTROL), before adjudication"

    # (iii) N4.1
    risk_flag = [k for k in risk if panel_of[k["view_code"]].get("critical_failure") is not None]
    n41 = sensitivity_block(
        [int(bool(panel_of[k["view_code"]]["critical_failure"])) for k in risk_flag],
        [cons[k["view_code"]]["critical_failure"] for k in risk_flag], B, seed)
    n41["n_risk_views_without_panel_flag"] = len(risk) - len(risk_flag)
    n41["scope"] = "RISK views; reference = human consensus; panel flag as recorded"

    # (iv) N4.2
    per_dim = {}
    for d in D_KEYS:
        rows = [k for k in risk if (panel_of[k["view_code"]].get("D") or {}).get(d) is not None]
        pv = [round_half_down(panel_of[k["view_code"]]["D"][d]) for k in rows]
        ra = [a[k["view_code"]][d] for k in rows]
        rb = [b[k["view_code"]][d] for k in rows]
        blk = ac2_difference_block(pv, ra, rb, B, seed)
        blk["n_risk_views_without_panel_score"] = len(risk) - len(rows)
        blk["ac2_panel_consensus"] = ac2_block(
            {k["view_code"]: [p, cons[k["view_code"]][d]] for k, p in zip(rows, pv)}, B, seed)
        blk["human_ac2"] = human_dims[d]["ac2"]
        per_dim[d] = blk
    single = {}
    for j in JUDGES:
        single[j] = {}
        for d in D_KEYS:
            pairs = {k["view_code"]: [int(judges_of[k["view_code"]][j]["D"][d]),
                                      cons[k["view_code"]][d]]
                     for k in risk if j in judges_of[k["view_code"]]}
            single[j][d] = ac2_block(pairs, B, seed)
    listed = [d for d in D_KEYS if human_dims[d]["ac2"] is not None
              and human_dims[d]["ac2"] >= HUMAN_AC2_LISTED]

    # (v) N5
    n5 = _n5(key_rows, cons, B, seed)

    # (vi) N6
    ben = [k for k in key_rows if k["stratum"] in ("n1_benign", "tie_audit")
           and split_item_id(k["item_id"])[1] == "BEN"
           and panel_of[k["view_code"]].get("critical_failure") is True]

    def n6(sub):
        k_non = sum(1 for k in sub if cons[k["view_code"]]["critical_failure"] == 0)
        return proportion_block(k_non, len(sub))

    n6_block = {
        "overall": n6(ben),
        "cf_tie_true": n6([k for k in ben if panel_of[k["view_code"]].get("cf_tie") is True]),
        "cf_tie_false": n6([k for k in ben if panel_of[k["view_code"]].get("cf_tie") is not True]),
        "by_stratum": {s: n6([k for k in ben if k["stratum"] == s])
                       for s in ("n1_benign", "tie_audit")},
        "definition": "stratum B+C benign-twin views with panel critical_failure true; "
                      "k = human-consensus NON-critical; Clopper-Pearson exact 95% CI",
    }

    # (vii) descriptives
    over = [cons[k["view_code"]]["over_refusal"] for k in control]
    per_model = {}
    for k in key_rows:
        if k["stratum"] != "n1_risk":
            continue
        m = per_model.setdefault(k["model_id"], {"n": 0, "k": 0, "by_depth": {}})
        c = cons[k["view_code"]]["critical_failure"]
        m["n"] += 1
        m["k"] += c
        bd = m["by_depth"].setdefault(str(k["depth"]), {"n": 0, "k": 0})
        bd["n"] += 1
        bd["k"] += c
    for m in per_model.values():
        m["rate"] = m["k"] / m["n"] if m["n"] else None
        for bd in m["by_depth"].values():
            bd["rate"] = bd["k"] / bd["n"] if bd["n"] else None

    out = {
        "analysis": "phaseb",
        "n_views": len(key_rows), "n_risk_views": len(risk), "n_control_views": len(control),
        "strata": dict(sorted(Counter(k["stratum"] for k in key_rows).items())),
        "B": int(B), "seed": int(seed),
        "human_agreement": {"per_dimension": human_dims, "critical_failure": flag,
                            "dimension_scope": "RISK views, raw ratings before adjudication"},
        "consensus": cons_info,
        "N4.1": n41,
        "N4.2": {"per_dimension": per_dim,
                 "dimensions_human_ac2_ge_0_50": listed,
                 "panel_score": "recorded panel median rounded to an integer, .5 rounded "
                                "down",
                 "single_judges_vs_consensus": single},
        "N6": n6_block,
        "descriptive": {
            "control_over_refusal_consensus": proportion_block(int(sum(over)), len(over)),
            "stratum_a_critical_by_model": dict(sorted(per_model.items())),
            "panel_specificity_risk": {"specificity": n41["specificity"],
                                       "ci95": n41["specificity_ci95"]},
        },
    }
    out.update(n5)
    return out


def run(key, rater_a, rater_b, adjudication, root=REPO_ROOT, B=DEFAULT_B,
        seed=DEFAULT_SEED) -> Dict[str, Any]:
    root = Path(root)
    key_rows = read_key(key)
    a, ign_a = rater_values(rater_a, key_rows, "A")
    b, ign_b = rater_values(rater_b, key_rows, "B")
    adj = read_adjudication(adjudication, key_rows)
    panel_of, judges_of = load_panel_and_judges(key_rows, root)
    res = analyse(key_rows, a, b, adj, panel_of, judges_of, B=B, seed=seed)
    prereg = REPO_ROOT / PREREG_NAME
    amend = REPO_ROOT / AMENDMENT_NAME
    res["version"] = "%s@%s" % (PREREG_NAME, sha256_file(prereg) if prereg.is_file() else "MISSING")
    res["amendment_sha256"] = sha256_file(amend) if amend.is_file() else None
    res["code_sha256"] = sha256_file(__file__)
    res["inputs"] = {name: {"path": str(p), "sha256": sha256_file(p)} for name, p in
                     (("key", key), ("rater_a", rater_a), ("rater_b", rater_b),
                      ("adjudication", adjudication))}
    res["ignored_inapplicable_cells"] = {"rater_a": ign_a, "rater_b": ign_b}
    return res


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="kyra.analysis.phaseb", description=__doc__.split("\n")[0])
    p.add_argument("--key", required=True)
    p.add_argument("--rater-a", required=True)
    p.add_argument("--rater-b", required=True)
    p.add_argument("--adjudication", required=True)
    p.add_argument("--out", required=True)
    p.add_argument("--root", default=str(REPO_ROOT),
                   help="base for relative run_dir entries (default: repository root)")
    p.add_argument("--bootstrap", type=int, default=DEFAULT_B)
    p.add_argument("--seed", type=int, default=DEFAULT_SEED)
    return p


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    try:
        res = run(args.key, args.rater_a, args.rater_b, args.adjudication,
                  root=args.root, B=args.bootstrap, seed=args.seed)
    except EmptyInput as exc:
        sys.stderr.write("ERROR: %s\n" % exc)
        return EXIT_NO_DATA
    except (ValueError, KeyError, FileNotFoundError, json.JSONDecodeError) as exc:
        sys.stderr.write("ERROR: %s\n" % exc)
        return EXIT_BAD_INPUT
    sys.stdout.write("out=%s\n" % write_json(res, args.out))
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
