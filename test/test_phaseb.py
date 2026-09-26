"""Tests for phase B: sampler outputs, rating workbooks, kyra.analysis.phaseb.

Analysis tests use SYNTHETIC fixtures only (fake run dirs, key and rater files
built here with known answers); no human label exists.  The committed sample,
key and rater workbooks under research/phaseB/ are checked structurally and
scanned for identity leaks.

Run:  cd /home/wjeong/adol && PYTHONPATH=src ~/envs/jeongwoncheol_adol/bin/python \
        -m unittest discover -s test -p "test_phaseb.py" -v
"""

from __future__ import annotations

import csv
import importlib.util
import io
import json
import math
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
import zlib
from collections import Counter
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "src"))

from kyra.analysis import phaseb  # noqa: E402

PB = REPO / "research" / "phaseB"
B_TEST = 40
D_KEYS = phaseb.D_KEYS


def _load_make_workbooks():
    spec = importlib.util.spec_from_file_location("make_workbooks", PB / "make_workbooks.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# -- synthetic fixture ---------------------------------------------------------

MODELS = (("fake/model-a", "ma"), ("fake/model-b", "mb"))
SCEN = range(1, 7)


def _crc(*parts) -> int:
    return zlib.crc32("|".join(str(p) for p in parts).encode())


def _panel_flag(item, depth):
    s = int(item[5])
    if item.endswith("-LOC") and item.startswith("R1"):
        return (depth == 6 and s <= 4) or (depth == 1 and s == 1), False
    if item.endswith("-BEN"):
        if depth == 6 and s in (1, 2):
            return True, True
        if depth == 3 and s in (3, 4):
            return True, False
        return False, False
    return s == 1, False  # R2-S0x-LOC singles


def _panel_d(run, item, depth, d):
    return (_crc(run, item, depth, d) % 5) / 2.0  # 0, .5, 1, 1.5, 2


def build_fixture(root: Path):
    """Two fake runs; returns (key_path, key_rows, truth) with truth = consensus by design."""
    items = {}
    for s in SCEN:
        items["R1-M0%d-LOC" % s] = 6
        items["R1-M0%d-BEN" % s] = 6
    for i in (1, 2, 3):
        items["R2-S0%d-LOC" % i] = 1
    key_rows = []
    for model, short in MODELS:
        run = root / "runs" / short
        run.mkdir(parents=True)
        panel, judges, resp = [], {j: [] for j in phaseb.JUDGES}, []
        for item, n in sorted(items.items()):
            for t in range(n):
                resp.append({"item_id": item, "condition": "base", "turn_index": t,
                             "user_text": "질문 %d" % (t + 1), "model_text": "답변 %d" % (t + 1)})
            for depth in range(1, n + 1):
                flag, tie = _panel_flag(item, depth)
                dvals = {d: _panel_d(short, item, depth, d) for d in D_KEYS}
                ctrl = item.endswith("-BEN")
                panel.append({"item_id": item, "condition": "base", "depth": depth,
                              "status": "ok", "D": dvals, "critical_failure": flag,
                              "cf_tie": tie, "is_control": ctrl,
                              "over_refusal": False if ctrl else None,
                              "helpfulness": 2 if ctrl else None})
                for j in phaseb.JUDGES:
                    err = (j == "J3" and item == "R2-S02-LOC")
                    rec = None if err else {
                        "D": {d: (math.floor(v) if j != "J2" else math.ceil(v))
                              for d, v in dvals.items()},
                        "critical_failure": flag}
                    judges[j].append({"item_id": item, "condition": "base", "depth": depth,
                                      "status": "ERROR" if err else "ok", "record": rec})
        for name, recs in [("panel.jsonl", panel), ("responses.jsonl", resp)] + [
                ("judge_%s.jsonl" % j, r) for j, r in judges.items()]:
            (run / name).write_text("".join(json.dumps(r, ensure_ascii=False) + "\n"
                                            for r in recs), encoding="utf-8")
        for s in SCEN:
            for depth in (1, 6):
                key_rows.append((model, run, "R1-M0%d-LOC" % s, depth, "n1_risk", "RISK"))
                key_rows.append((model, run, "R1-M0%d-BEN" % s, depth, "n1_benign", "CONTROL"))
        for s in (3, 4):
            key_rows.append((model, run, "R1-M0%d-BEN" % s, 3, "tie_audit", "CONTROL"))
        for i in (1, 2, 3):
            key_rows.append((model, run, "R2-S0%d-LOC" % i, 1, "calibration_spread", "RISK"))
    rows = []
    for n, (model, run, item, depth, stratum, task) in enumerate(key_rows, start=1):
        flag, tie = _panel_flag(item, depth)
        rows.append({"view_code": "V%03d" % n, "stratum": stratum, "run_dir": str(run),
                     "model_id": model, "item_id": item, "condition": "base",
                     "depth": depth, "task": task,
                     "panel_critical_failure": "true" if flag else "false",
                     "panel_cf_tie": "true" if tie else "false"})
    key = root / "key.csv"
    with key.open("w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    truth = {}
    for r in rows:
        flag, _ = _panel_flag(r["item_id"], r["depth"])
        short = Path(r["run_dir"]).name
        t = {d: phaseb.round_half_down(_panel_d(short, r["item_id"], r["depth"], d))
             for d in D_KEYS}
        t["critical_failure"] = int(flag)
        if r["item_id"].endswith("-BEN") and flag:
            # tie audit truth: only model-a R1-M01 depth 6 is really critical
            t["critical_failure"] = int(short == "ma" and r["item_id"] == "R1-M01-BEN")
        t["over_refusal"] = int(_crc(r["view_code"]) % 4 == 0)
        t["helpfulness"] = _crc(r["view_code"], "h") % 3
        truth[r["view_code"]] = t
    return key, rows, truth


def write_rater_csv(path: Path, rows, values):
    header = [h for _, h in phaseb.SHEET_COLUMNS if _ != "transcript"]
    keys = [k for k, _ in phaseb.SHEET_COLUMNS if k != "transcript"]
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        for r in rows:
            v = values[r["view_code"]]
            fields = phaseb.fields_for(r["task"])
            out = []
            for k in keys:
                if k == "view_code":
                    out.append(r["view_code"])
                elif k == "task":
                    out.append(r["task"])
                elif k in fields:
                    x = v[k]
                    out.append((phaseb.YES if x else phaseb.NO) if k in phaseb.BINARY_FIELDS
                               else str(x))
                else:
                    out.append("")
            w.writerow(out)


def write_adj(path: Path, rows):
    with path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["view_code", "field", "value", "rationale"])
        w.writerows(rows)


def run_cli(args):
    env = {"PYTHONPATH": str(REPO / "src"), "PATH": "/usr/bin:/bin"}
    return subprocess.run([sys.executable, "-m", "kyra.analysis.phaseb"] + args,
                          cwd=str(REPO), env=env, capture_output=True, text=True)


class FixtureCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.key, self.rows, self.truth = build_fixture(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def raters(self, a_vals, b_vals, adj_rows=()):
        a, b, adj = self.root / "a.csv", self.root / "b.csv", self.root / "adj.csv"
        write_rater_csv(a, self.rows, a_vals)
        write_rater_csv(b, self.rows, b_vals)
        write_adj(adj, adj_rows)
        return a, b, adj

    def analyse(self, a_vals, b_vals, adj_rows=()):
        a, b, adj = self.raters(a_vals, b_vals, adj_rows)
        return phaseb.run(self.key, a, b, adj, root=self.root, B=B_TEST, seed=20260926)


# -- analysis tests --------------------------------------------------------------


class TestPhaseBAnalysis(FixtureCase):
    def test_identical_raters_give_ac2_one_and_kappa_one(self):
        res = self.analyse(self.truth, self.truth)
        for d in D_KEYS:
            self.assertAlmostEqual(res["human_agreement"]["per_dimension"][d]["ac2"], 1.0)
        flag = res["human_agreement"]["critical_failure"]
        self.assertEqual(flag["raw_agreement"], 1.0)
        self.assertAlmostEqual(flag["cohen_kappa"], 1.0)
        self.assertEqual(res["consensus"]["n_adjudicated"], 0)
        self.assertEqual(res["n_views"], len(self.rows))

    def test_panel_equal_to_humans_gives_zero_difference_and_sensitivity_one(self):
        res = self.analyse(self.truth, self.truth)
        for d in D_KEYS:
            blk = res["N4.2"]["per_dimension"][d]
            self.assertAlmostEqual(blk["difference"], 0.0)
            self.assertAlmostEqual(blk["ac2_panel_consensus"]["ac2"], 1.0)
        self.assertEqual(res["N4.1"]["sensitivity"], 1.0)
        self.assertEqual(res["N4.1"]["specificity"], 1.0)
        self.assertEqual(res["N4.2"]["dimensions_human_ac2_ge_0_50"], list(D_KEYS))
        # J3 has one ERROR view per run -> 2 fewer subjects than J1
        n1 = res["N4.2"]["single_judges_vs_consensus"]["J1"]["D1"]["n_subjects"]
        n3 = res["N4.2"]["single_judges_vs_consensus"]["J3"]["D1"]["n_subjects"]
        self.assertEqual(n1 - n3, 2)

    def test_raters_flipping_every_flag_give_sensitivity_zero(self):
        flipped = {c: dict(v, critical_failure=1 - int(
            self.rows[int(c[1:]) - 1]["panel_critical_failure"] == "true"))
            for c, v in self.truth.items()}
        res = self.analyse(flipped, flipped)
        self.assertEqual(res["N4.1"]["sensitivity"], 0.0)
        self.assertEqual(res["N4.1"]["specificity"], 0.0)
        self.assertEqual(res["N4.1"]["true_positive"], 0)

    def test_n6_known_counts_and_clopper_pearson(self):
        res = self.analyse(self.truth, self.truth)
        ov = res["N6"]["overall"]
        self.assertEqual((ov["n"], ov["k"]), (8, 7))
        lo, hi = ov["ci95_clopper_pearson"]
        from scipy import stats
        self.assertAlmostEqual(lo, stats.beta.ppf(0.025, 7, 2))
        self.assertAlmostEqual(hi, stats.beta.ppf(0.975, 8, 1))
        self.assertEqual((res["N6"]["cf_tie_true"]["n"], res["N6"]["cf_tie_true"]["k"]), (4, 3))
        self.assertEqual((res["N6"]["cf_tie_false"]["n"], res["N6"]["cf_tie_false"]["k"]), (4, 4))
        self.assertEqual(phaseb.clopper_pearson(5, 5)[1], 1.0)
        self.assertAlmostEqual(phaseb.clopper_pearson(5, 5)[0], 0.025 ** (1 / 5))

    def test_n5_counts_and_estimates_present(self):
        res = self.analyse(self.truth, self.truth)
        c = res["N5.1"]["counts"]
        self.assertEqual((c["n_conversations"], c["depth1_critical"], c["depth6_critical"]),
                         (12, 2, 8))
        bc = res["N5.2"]["benign_counts"]
        self.assertEqual((bc["depth1_critical"], bc["depth6_critical"]), (0, 1))
        self.assertIsNone(res["N5.1"]["error"])
        self.assertGreater(res["N5.1"]["or"], 1.0)
        self.assertIn("interaction_p_holm", res["N5.2"])
        a = res["descriptive"]["stratum_a_critical_by_model"]["fake/model-a"]
        self.assertEqual((a["n"], a["k"]), (12, 5))

    def test_disagreement_uses_adjudicated_value(self):
        b = {c: dict(v) for c, v in self.truth.items()}
        risk = next(r["view_code"] for r in self.rows if r["task"] == "RISK")
        b[risk]["D2"] = (self.truth[risk]["D2"] + 1) % 3
        a_p, b_p, adj = self.raters(self.truth, b,
                                    [[risk, "D2", str(b[risk]["D2"]), "joint reading"]])
        key_rows = phaseb.read_key(self.key)
        av, _ = phaseb.rater_values(a_p, key_rows, "A")
        bv, _ = phaseb.rater_values(b_p, key_rows, "B")
        cons, info = phaseb.build_consensus(key_rows, av, bv,
                                            phaseb.read_adjudication(adj, key_rows))
        self.assertEqual(cons[risk]["D2"], b[risk]["D2"])
        self.assertNotEqual(cons[risk]["D2"], self.truth[risk]["D2"])
        self.assertEqual(info["n_adjudicated"], 1)

    def test_missing_adjudication_is_refused_by_cli(self):
        b = {c: dict(v) for c, v in self.truth.items()}
        risk = next(r["view_code"] for r in self.rows if r["task"] == "RISK")
        b[risk]["critical_failure"] = 1 - b[risk]["critical_failure"]
        a_p, b_p, adj = self.raters(self.truth, b)
        out = self.root / "out.json"
        p = run_cli(["--key", str(self.key), "--rater-a", str(a_p), "--rater-b", str(b_p),
                     "--adjudication", str(adj), "--out", str(out),
                     "--root", str(self.root), "--bootstrap", str(B_TEST)])
        self.assertEqual(p.returncode, 4, p.stderr)
        self.assertIn("no adjudication row", p.stderr)
        self.assertFalse(out.exists())
        # the same input with the adjudication row runs
        write_adj(adj, [[risk, "critical_failure", "예", "joint reading"]])
        p = run_cli(["--key", str(self.key), "--rater-a", str(a_p), "--rater-b", str(b_p),
                     "--adjudication", str(adj), "--out", str(out),
                     "--root", str(self.root), "--bootstrap", str(B_TEST)])
        self.assertEqual(p.returncode, 0, p.stderr)
        res = json.loads(out.read_text(encoding="utf-8"))
        for claim in ("N4.1", "N4.2", "N5.1", "N5.2", "N6"):
            self.assertIn(claim, res)
        self.assertNotIn("PASS", json.dumps(res))
        self.assertNotIn("FAIL", json.dumps(res))

    def test_bad_adjudication_rows_are_refused(self):
        risk = next(r["view_code"] for r in self.rows if r["task"] == "RISK")
        ctrl = next(r["view_code"] for r in self.rows if r["task"] == "CONTROL")
        agreed = self.truth[risk]["D1"]
        cases = [
            [[risk, "D1", str((agreed + 1) % 3), "x"]],          # contradicts agreement
            [[risk, "D1", str(agreed), ""]],                     # empty rationale
            [[ctrl, "D1", "1", "x"]],                            # field not for CONTROL
            [["V999", "D1", "1", "x"]],                          # unknown view
            [[risk, "D1", str(agreed), "x"], [risk, "D1", str(agreed), "y"]],  # duplicate
        ]
        for rows in cases:
            with self.subTest(rows=rows):
                with self.assertRaises(phaseb.PhaseBError):
                    self.analyse(self.truth, self.truth, rows)

    def test_missing_required_cell_is_refused(self):
        a = {c: dict(v) for c, v in self.truth.items()}
        a_p, b_p, adj = self.raters(a, self.truth)
        text = a_p.read_text(encoding="utf-8").splitlines()
        cells = text[1].split(",")
        cells[2] = ""  # D1 of the first row (or over_refusal slot -> still required/empty)
        risk_first = self.rows[0]["task"] == "RISK"
        self.assertTrue(risk_first)
        text[1] = ",".join(cells)
        a_p.write_text("\n".join(text) + "\n", encoding="utf-8")
        with self.assertRaises(phaseb.PhaseBError):
            phaseb.run(self.key, a_p, b_p, adj, root=self.root, B=B_TEST)

    def test_round_half_down(self):
        cases = {0: 0, 0.5: 0, 1: 1, 1.5: 1, 2: 2, 1.0: 1}
        for x, want in cases.items():
            self.assertEqual(phaseb.round_half_down(x), want)
        self.assertIsNone(phaseb.round_half_down(None))

    def test_xlsx_roundtrip_matches_csv(self):
        mw = _load_make_workbooks()
        texts = mw.transcripts(self.rows, root=self.root)
        path = self.root / "rater_x.xlsx"
        mw.build_workbook(path, list(reversed(self.rows)), texts, "X")
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb[phaseb.RATER_SHEET]
        col = {c.value.split()[0]: c.column for c in ws[1]}
        greyed_written = 0
        for row in range(2, ws.max_row + 1):
            code = ws.cell(row=row, column=col["view_code"]).value
            task = ws.cell(row=row, column=col["task"]).value
            for f in phaseb.fields_for(task):
                v = self.truth[code][f]
                ws.cell(row=row, column=col[f]).value = (
                    (phaseb.YES if v else phaseb.NO) if f in phaseb.BINARY_FIELDS else v)
            if task == "RISK" and greyed_written == 0:
                ws.cell(row=row, column=col["helpfulness"]).value = 2  # greyed cell
                greyed_written = 1
        wb.save(path)
        key_rows = phaseb.read_key(self.key)
        xv, ignored = phaseb.rater_values(path, key_rows, "X")
        self.assertEqual(ignored, 1)
        for code, vals in xv.items():
            for f, v in vals.items():
                self.assertEqual(v, self.truth[code][f], (code, f))
        # transcript shows turns 1..depth only
        wb = load_workbook(path, read_only=True)
        for r in wb[phaseb.RATER_SHEET].iter_rows(min_row=2, values_only=True):
            depth = next(k["depth"] for k in key_rows if k["view_code"] == r[0])
            marks = re.findall(r"^\[(\d+)\] AI:", r[2], re.MULTILINE)
            self.assertEqual([int(m) for m in marks], list(range(1, depth + 1)))

    def test_redaction_rule(self):
        mw = _load_make_workbooks()
        masked = ["저는 구글에서 만든 대규모 언어 모델입니다.", "저는 Qwen입니다.",
                  "LG AI Research에서 개발한 EXAONE", "Google의 데이터 센터에 있는 서버",
                  "Alibaba Cloud에서 제공하는 AI 어시스턴트"]
        kept = ["카카오톡으로 보내 봐", "네이버 사전을 써 봐", "구글 드라이브에 올려",
                "청소년사이버상담센터 1388/카카오 채널", "네이버에서 제공하는 상담 안내"]
        for s in masked:
            self.assertGreater(mw.redact(s)[1], 0, s)
        for s in kept:
            self.assertEqual(mw.redact(s), (s, 0), s)


# -- committed phase-B files ----------------------------------------------------------


def _read_csv(path):
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


class TestPhaseBSample(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.sample = _read_csv(PB / "sample_phaseB_v1.csv")
        cls.key = _read_csv(PB / "key_phaseB_v1.csv")
        cls.calib = _read_csv(PB / "calibration_phaseB_v1.csv")
        cls.summary = json.loads((PB / "sample_phaseB_v1_summary.json").read_text("utf-8"))

    def test_columns_and_counts(self):
        want = ["view_code", "stratum", "run_dir", "model_id", "item_id", "condition",
                "depth", "task", "panel_critical_failure", "panel_cf_tie"]
        self.assertEqual(list(self.sample[0]), want)
        self.assertEqual(len(self.sample), 300)
        self.assertEqual([r["view_code"] for r in self.sample],
                         ["V%03d" % i for i in range(1, 301)])
        c = Counter(r["stratum"] for r in self.sample)
        self.assertEqual((c["n1_risk"], c["n1_benign"]), (120, 120))
        self.assertEqual(c["tie_audit"] + c["calibration_spread"], 60)
        self.assertEqual(len(self.calib), 10)
        for s, k in zip(self.sample, self.key):
            self.assertEqual({f: s[f] for f in want}, {f: k[f] for f in want})

    def test_stratum_a_design(self):
        a = [r for r in self.sample if r["stratum"] == "n1_risk"]
        pairs = {(r["item_id"], r["model_id"]) for r in a}
        self.assertEqual(len(pairs), 60)
        self.assertEqual(set(Counter(m for _, m in pairs).values()), {12})
        self.assertLessEqual(set(Counter(s for s, _ in pairs).values()), {2, 3})
        self.assertEqual(len({s for s, _ in pairs}), 24)
        for s, m in pairs:
            self.assertEqual(sorted(int(r["depth"]) for r in a
                                    if (r["item_id"], r["model_id"]) == (s, m)), [1, 6])
        b = {(r["item_id"], r["model_id"], r["depth"]) for r in self.sample
             if r["stratum"] == "n1_benign"}
        want = {(s[:-4] + "-BEN", m, d) for s, m in pairs for d in ("1", "6")}
        self.assertEqual(b, want)

    def test_stratum_c_and_d(self):
        b_flag = [r for r in self.sample if r["stratum"] == "n1_benign"
                  and r["panel_critical_failure"] == "true"]
        c = [r for r in self.sample if r["stratum"] == "tie_audit"]
        self.assertEqual(len(c) + len(b_flag),
                         self.summary["stratum_c"]["ben_views_panel_flagged_total"])
        self.assertTrue(all(r["item_id"].endswith("-BEN") and r["panel_critical_failure"]
                            == "true" for r in c))
        a_views = {(r["run_dir"], r["item_id"], r["depth"]) for r in self.sample
                   if r["stratum"] == "n1_risk"}
        d = [r for r in self.key if r["stratum"] == "calibration_spread"]
        self.assertTrue(all(r["task"] == "RISK" for r in d))
        self.assertFalse({(r["run_dir"], r["item_id"], r["depth"]) for r in d} & a_views)
        per_model = Counter(r["model_id"] for r in d)
        self.assertLessEqual(max(per_model.values()) - min(per_model.values()), 1)
        halves = Counter(r["d_half"] for r in d)
        self.assertLessEqual(abs(halves["disagreement"] - halves["random"]), 1)
        self.assertTrue(all(r["judge_disagreement"] == "true" for r in d
                            if r["d_half"] == "disagreement"))

    def test_calibration_outside_sample(self):
        conv = {(r["run_dir"], r["item_id"]) for r in self.sample}
        for r in self.calib:
            self.assertEqual(r["task"], "RISK")
            self.assertNotIn((r["run_dir"], r["item_id"]), conv)
        self.assertEqual(set(Counter(r["model_id"] for r in self.calib).values()), {2})


class TestRaterWorkbooks(unittest.TestCase):
    FILES = ("rater_A_v1.xlsx", "rater_B_v1.xlsx", "calibration_v1.xlsx")

    @classmethod
    def setUpClass(cls):
        key = _read_csv(PB / "key_phaseB_v1.csv") + _read_csv(PB / "calibration_phaseB_v1.csv")
        terms = {"panel", "J1", "J2", "J3", "phaseA", "result/raw", "/main/",
                 "n1_risk", "n1_benign", "tie_audit", "calibration_spread",
                 "llama", "mistral", "phi-4", "exaone", "qwen", "gemma", "kanana",
                 "hyperclova", "엑사원", "큐웬", "젬마", "카나나", "하이퍼클로바"}
        for r in key:
            terms.add(r["model_id"])
            terms.add(r["run_dir"])
            terms.update(p for p in r["run_dir"].split("/") if p not in ("result", "raw", "main"))
            terms.add(r["item_id"])
            terms.add(r["model_id"].split("/")[-1])
        cls.terms = sorted(terms)
        cls.key = key

    def _hits(self, text):
        low = text.lower()
        return [t for t in self.terms
                if (t in text if re.fullmatch(r"J[123]", t) else t.lower() in low)]

    def test_no_identity_or_judge_output_in_any_cell(self):
        from openpyxl import load_workbook
        n_cells = 0
        for f in self.FILES:
            wb = load_workbook(PB / "raters" / f)
            for ws in wb.worksheets:
                self.assertEqual(self._hits(ws.title), [])
                for row in ws.iter_rows():
                    for c in row:
                        if c.value is None:
                            continue
                        n_cells += 1
                        self.assertEqual(self._hits(str(c.value)), [], (f, ws.title, c.coordinate))
            props = " ".join(str(v) for v in vars(wb.properties).values() if v)
            self.assertEqual(self._hits(props), [], f)
        self.assertGreater(n_cells, 600)

    def test_no_identity_in_raw_xlsx_parts(self):
        long_terms = [t for t in self.terms if len(t) >= 6 and not re.fullmatch(r"J[123]", t)]
        for f in self.FILES:
            with zipfile.ZipFile(PB / "raters" / f) as z:
                blob = b"".join(z.read(n) for n in z.namelist()).decode("utf-8", "replace")
            low = blob.lower()
            self.assertEqual([t for t in long_terms if t.lower() in low], [], f)

    def test_scan_catches_a_planted_model_id(self):
        self.assertTrue(self._hits("대화 ... Qwen/Qwen2.5-14B-Instruct ..."))
        self.assertTrue(self._hits("J2"))

    def test_workbooks_open_with_dropdowns_and_all_views(self):
        from openpyxl import load_workbook
        key = _read_csv(PB / "key_phaseB_v1.csv")
        want = {"rater_A_v1.xlsx": sorted(r["view_code"] for r in key),
                "rater_B_v1.xlsx": sorted(r["view_code"] for r in key),
                "calibration_v1.xlsx": ["C%02d" % i for i in range(1, 11)]}
        orders = {}
        for f in self.FILES:
            wb = load_workbook(PB / "raters" / f)
            self.assertEqual(wb.sheetnames, [phaseb.GUIDE_SHEET, phaseb.RATER_SHEET])
            ws = wb[phaseb.RATER_SHEET]
            header = [c.value.split()[0] for c in ws[1]]
            self.assertEqual(header, [k for k, _ in phaseb.SHEET_COLUMNS])
            codes = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
            self.assertEqual(sorted(codes), want[f])
            orders[f] = codes
            dv = {d.formula1: d for d in ws.data_validations.dataValidation}
            self.assertIn('"0,1,2"', dv)
            self.assertIn('"%s,%s"' % (phaseb.YES, phaseb.NO), dv)
            n_risk = sum(1 for r in range(2, ws.max_row + 1)
                         if ws.cell(row=r, column=2).value == "RISK")
            n_ctrl = len(codes) - n_risk
            n_score = sum(len(list(rg.cells)) for rg in dv['"0,1,2"'].sqref.ranges)
            n_yn = sum(len(list(rg.cells)) for rg in dv['"%s,%s"' % (phaseb.YES, phaseb.NO)]
                       .sqref.ranges)
            self.assertEqual(n_score, 6 * n_risk + n_ctrl)
            self.assertEqual(n_yn, len(codes) + n_ctrl)
            guide = " ".join(str(c.value) for c in wb[phaseb.GUIDE_SHEET]["A"] if c.value)
            for word in ("D1", "D6", "CF1", "CF6", "40", "독립", "블라인딩"):
                self.assertIn(word, guide)
        self.assertNotEqual(orders["rater_A_v1.xlsx"], orders["rater_B_v1.xlsx"])


if __name__ == "__main__":
    unittest.main()
